from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from datetime import date
from pathlib import Path

OUT = Path("docs/QuickType_Technical_Architecture.docx")
OUT.parent.mkdir(exist_ok=True)

BLUE = "2E74B5"
DARK = "1F4D78"
INK = "243746"
MUTED = "64748B"
PALE = "E8EEF5"
LIGHT = "F2F4F7"
WHITE = "FFFFFF"
RED = "9B1C1C"

doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"; normal.font.size = Pt(11); normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.25
for name, size, color, before, after in [
    ("Heading 1",16,BLUE,18,10),("Heading 2",13,BLUE,14,7),("Heading 3",12,DARK,10,5)]:
    s=styles[name]; s.font.name="Calibri"; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color)
    s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True
styles["Title"].font.name="Calibri"; styles["Title"].font.size=Pt(30); styles["Title"].font.bold=True; styles["Title"].font.color.rgb=RGBColor.from_string(DARK)
styles["Subtitle"].font.name="Calibri"; styles["Subtitle"].font.size=Pt(15); styles["Subtitle"].font.color.rgb=RGBColor.from_string(MUTED)

code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
code_style.font.name="Consolas"; code_style.font.size=Pt(8.5); code_style.font.color.rgb=RGBColor.from_string("172033")
code_style.paragraph_format.space_before=Pt(3); code_style.paragraph_format.space_after=Pt(7); code_style.paragraph_format.left_indent=Inches(.18); code_style.paragraph_format.right_indent=Inches(.18)

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn("w:shd"))
    if shd is None: shd=OxmlElement("w:shd"); tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def margins(cell, top=80, start=120, bottom=80, end=120):
    tc=cell._tc.get_or_add_tcPr(); m=tc.first_child_found_in("w:tcMar")
    if m is None: m=OxmlElement("w:tcMar"); tc.append(m)
    for tag,val in (("top",top),("start",start),("bottom",bottom),("end",end)):
        el=m.find(qn("w:"+tag))
        if el is None: el=OxmlElement("w:"+tag); m.append(el)
        el.set(qn("w:w"),str(val)); el.set(qn("w:type"),"dxa")

def set_repeat(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement("w:tblHeader"); el.set(qn("w:val"),"true"); trPr.append(el)

def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.LEFT; t.autofit=False; t.style="Table Grid"
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,PALE); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(c)
        for r in c.paragraphs[0].runs: r.bold=True; r.font.color.rgb=RGBColor.from_string(DARK); r.font.size=Pt(9)
    set_repeat(t.rows[0])
    for vals in rows:
        cells=t.add_row().cells
        for i,v in enumerate(vals):
            cells[i].text=str(v); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(cells[i])
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(1); p.paragraph_format.line_spacing=1.08
                for r in p.runs: r.font.size=Pt(8.7)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def bullet(text, level=0):
    p=doc.add_paragraph(style="List Bullet" if level==0 else "List Bullet 2"); p.add_run(text)
    p.paragraph_format.left_indent=Inches(.375+.25*level); p.paragraph_format.first_line_indent=Inches(-.188); p.paragraph_format.space_after=Pt(4); p.paragraph_format.line_spacing=1.25
    return p

def number(text):
    p=doc.add_paragraph(style="List Number"); p.add_run(text); p.paragraph_format.left_indent=Inches(.375); p.paragraph_format.first_line_indent=Inches(-.188); p.paragraph_format.space_after=Pt(4); return p

def code(text):
    p=doc.add_paragraph(style="Code Block"); p.add_run(text)
    pPr=p._p.get_or_add_pPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),LIGHT); pPr.append(shd)
    return p

def note(label,text):
    t=doc.add_table(rows=1,cols=1); t.autofit=False; t.columns[0].width=Inches(6.5); c=t.cell(0,0); shade(c,LIGHT); margins(c,120,160,120,160)
    p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); r=p.add_run(label+": "); r.bold=True; r.font.color.rgb=RGBColor.from_string(DARK); p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

def page_break(): doc.add_page_break()

# Header/footer
hp=sec.header.paragraphs[0]; hp.text="QUICKTYPE  |  TECHNICAL ARCHITECTURE"; hp.alignment=WD_ALIGN_PARAGRAPH.LEFT
for r in hp.runs: r.font.size=Pt(8); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(MUTED)
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fp.add_run("QuickType repository guide  •  ")
field=OxmlElement("w:fldSimple"); field.set(qn("w:instr"),"PAGE"); fp._p.append(field)
for r in fp.runs: r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(MUTED)

# Cover
doc.add_paragraph().paragraph_format.space_after=Pt(80)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("TECHNICAL REFERENCE GUIDE"); r.bold=True; r.font.size=Pt(10); r.font.color.rgb=RGBColor.from_string(BLUE)
p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("QuickType")
p=doc.add_paragraph(style="Subtitle"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("Browser configuration, embedded firmware, HID routing, and expansion execution")
doc.add_paragraph().paragraph_format.space_after=Pt(42)
note("Purpose","Explain how the project works from the browser UI down to serial framing, persistent configuration, dual-core USB-host processing, rule matching, RTC substitution, and laptop-facing HID reports.")
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(70)
r=p.add_run(f"Repository snapshot reviewed {date.today().isoformat()}"); r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string(MUTED)
page_break()

doc.add_heading("1. System at a glance",1)
doc.add_paragraph("QuickType is a hardware text-expansion and input-routing system. A single-file web configurator edits a JSON rule set and writes it over Web Serial to a primary RP2040. The primary appears to the computer as a composite USB device (CDC serial, keyboard, Consumer Control, and mouse), stores configuration in LittleFS, reads a DS3231 real-time clock, and transforms input from a keypad or keyboard into normal HID output.")
code("Browser (index.html)\n  │ Web Serial: newline-delimited JSON, 115200 baud\n  ▼\n+Primary RP2040 (RTC_Setup.ino) ── I²C GPIO6/7 ── DS3231 RTC\n  │ laptop-facing USB: CDC + keyboard + consumer + mouse\n  │\n  ├─ one-board input: Pico-PIO-USB on GPIO0/1\n  └─ two-board input: UART GPIO4/5, 115200 baud\n          ▲\n          │ framed binary HID descriptor/report packets\n+Bridge RP2040 (UsbUart_Bridge.ino) ── native USB host ── keypad/keyboard/receiver")
table(["Layer","Primary responsibility","Key source"],[
    ("Configurator","Edit, validate, import/export, discover devices, send commands","index.html"),
    ("Primary firmware","Persist/compile rules, match triggers, resolve tokens, emit HID","firmware/RTC_Setup/RTC_Setup.ino"),
    ("Bridge firmware","Native USB host; forward descriptors and raw reports over UART","firmware/UsbUart_Bridge/UsbUart_Bridge.ino"),
    ("Shared content","Versioned common expansion catalogs and sets","expansion-sets/*.json"),
    ("Diagnostics","Host/bridge probes and long-running telemetry monitors","firmware/*Diagnostic, tools/*"),
], [1.15,3.1,2.25])
doc.add_heading("Operational modes",2)
bullet("One-board mode: the primary RP2040 hosts the downstream keypad through Pico-PIO-USB on GPIO0/GPIO1.")
bullet("Two-board mode: the bridge hosts the USB device natively and forwards raw HID information over crossed GPIO4/GPIO5 UART wiring plus common ground.")
bullet("The primary automatically selects UART input after valid bridge heartbeats. If heartbeats stop for 1.5 seconds, it releases held input and falls back to PIO USB.")
bullet("Only the selected downstream input path is serviced, while laptop-facing USB remains attached.")

doc.add_heading("2. Repository map",1)
table(["Path","What to change there"],[
    ("index.html","All browser HTML, CSS, and JavaScript. There is no framework or bundler."),
    ("firmware/RTC_Setup/","Production primary-board sketch and its UART protocol header."),
    ("firmware/UsbUart_Bridge/","Production keyboard-facing bridge sketch and matching protocol header."),
    ("firmware/third_party/Pico_PIO_USB/","Pinned downstream USB-host implementation with project patch notes."),
    ("firmware/Host_Diagnostic/","Minimal PIO host test sketch."),
    ("firmware/Bridge_Diagnostic/","Native-host diagnostic sketch."),
    ("expansion-sets/","Catalog plus versioned reusable expansion-set JSON."),
    ("tools/","Serial and USB stability probes for macOS/Linux and Windows."),
], [2.05,4.45])
note("Build-stamp rule","Whenever index.html changes, its quicktype-build meta tag must be updated to the current local timestamp including UTC offset (repository AGENTS.md).")

page_break()
doc.add_heading("3. Browser configurator",1)
doc.add_paragraph("The website is intentionally monolithic: markup, responsive styling, application state, validation, serialization, serial I/O, and shared-set logic all live in index.html. initApp() caches elements, binds UI events, initializes theme/navigation/preferences/shared sets, validates Web Serial support, chooses the requested view, renders configuration state, and begins authorized-device discovery.")
doc.add_heading("Application state and editor model",2)
bullet("The in-memory configuration is the editing authority. Dirty-state markers distinguish editor changes from the last device read/write.")
bullet("A maximum of 128 active rules is shared between typed expansions and keypad mappings. Browser-side payload capacity is 30,000 bytes, leaving margin below the firmware’s 32 KB request ceiling.")
bullet("Keypad rules are identified from physical trigger names such as KP_0 and KP_ENTER; typed rules occupy stable R01…R128 slots.")
bullet("Normalization accepts current configuration plus legacy keymap shapes, repairs defaults, sanitizes placeholders, and converts older action types.")
bullet("Trigger validation rejects duplicates and prefix ambiguity when scopes overlap. This prevents one trigger from making a longer trigger unreachable.")
doc.add_heading("Browser persistence and external files",2)
table(["Data","Persistence/transport"],[
    ("Theme, navigation, auto-connect, shared-set subscriptions","Browser localStorage"),
    ("Device configuration","In memory until explicitly written over Web Serial"),
    ("Backup/transfer","Downloaded or imported JSON files"),
    ("Published common sets","Fetched from ./expansion-sets/catalog.json and referenced set URLs"),
], [2.15,4.35])
doc.add_heading("Connection lifecycle",2)
number("navigator.serial.getPorts() discovers already-authorized ports; requestPort() prompts for a new one using QuickType USB filters when possible.")
number("connectDevice() opens the port at 115200 baud, acquires reader/writer streams, starts the asynchronous read loop, and sends ping.")
number("The ping response supplies device identity, firmware version, configuration presence, active rule count, runtime enable states, selected input mode, and downstream-device status.")
number("The browser may read configuration, synchronize required common sets, and poll clock/telemetry. Hot-plug listeners keep the device list current.")
number("Disconnect drains or rejects outstanding request promises, releases stream locks, closes the port, and resets connection UI state.")

doc.add_heading("4. Web Serial protocol",1)
doc.add_paragraph("Each message is one minified JSON object terminated by a newline. Requests carry protocol version qt=1, a monotonically increasing integer id, a command, and command-specific fields. sendDeviceCommand() stores a pending promise keyed by id; handleSerialLine() validates and resolves or rejects it. Unrelated debug text is ignored because it does not satisfy the protocol envelope.")
code('{"qt":1,"id":42,"command":"get-clock"}\n{"qt":1,"id":42,"ok":true,"type":"clock","data":{...}}')
table(["Command","Effect","Persistence"],[
    ("ping","Return identity/capability/runtime state","None"),
    ("get-config","Stream saved JSON file inside a config response","Read only"),
    ("set-config","Validate, compile, atomically save, activate","LittleFS"),
    ("get-clock / set-clock","Read/write DS3231 and timezone metadata","RTC + metadata"),
    ("get-telemetry","Return counters, modes, mounts, heap and timing","Read only"),
    ("set-expansions-enabled","Pause/resume all configured rules","Runtime only"),
    ("set-keypad-expansions-enabled","Pause/resume configured physical-key actions","Runtime only"),
    ("factory-reset","Remove saved config and restore legacy behavior","LittleFS deletion"),
    ("reset-usb","Acknowledge, detach, and reattach laptop-facing USB","Runtime only"),
], [1.75,3.2,1.55])
note("Protocol boundary","Firmware caps the line/configuration size at 32 KB. Malformed JSON, wrong protocol versions, invalid state/clock/config values, missing configuration, and unknown commands return structured error codes.")

page_break()
doc.add_heading("5. Configuration data model",1)
doc.add_paragraph("The browser stores a friendly superset for editing, then buildDeviceConfiguration() removes inactive rule slots and serializes only active rules. The primary validates and compiles the JSON into fixed-capacity runtime structures so expansion execution does not parse the file repeatedly.")
table(["Concept","Representative fields","Semantics"],[
    ("Device","name, uniqueId, color","Presentation and stable device identity."),
    ("Rule","id, enabled, label, type, text, steps, keyDelay","Action definition and execution timing."),
    ("Trigger","trigger, triggerPattern, scope","Physical key or typed suffix; keyboard/numpad/any source scope."),
    ("Placeholder","name → value","Reusable literal substitution resolved when an action runs."),
    ("Managed set metadata","setId, entryId, revision/content identity","Tracks common rules without overwriting personal rules."),
], [1.25,2.35,2.9])
doc.add_heading("Rule kinds",2)
bullet("Physical key (trigger = key): intercepts a named key. Assigned actions execute; unassigned keys pass through unchanged.")
bullet("Typed trigger: buffers recent printable input and attempts suffix matching only when Space, Tab, Enter, or Escape arrives.")
bullet("Expansion/template: emits literal text with token and inline-keystroke interpretation.")
bullet("Shortcut: sends one HID chord; legacy shortcut rules are converted to double-brace syntax by the editor.")
bullet("Steps: execute multiple values with per-step delays; values can request expansion typing, placeholder resolution, key:<shortcut>, or literal template output.")
doc.add_heading("Typed-trigger delimiter behavior",2)
table(["Following key","On match","On no match"],[
    ("Space","Remove trigger, emit expansion, then restore Space","Forward normally"),
    ("Tab / Enter / Escape","Remove trigger, emit expansion, consume delimiter","Forward normally"),
    ("Other input","Do not expand; continue buffering/forwarding","Forward normally"),
], [1.35,2.8,2.35])
doc.add_paragraph("The scope check uses the recorded source for typed characters, allowing rules to be limited to the full keyboard, numeric keypad, or either. The hidden ;;; trigger followed by a delimiter emits a reference listing of active typed expansions and configured keypad actions.")

doc.add_heading("6. Safe persistence on the primary",1)
doc.add_paragraph("Configuration writes are transactional at the file level. The firmware first measures and compiles the candidate. During LittleFS work it suspends the downstream host stack and disables the watchdog while continuing to service laptop-facing native USB cooperatively.")
number("Serialize the candidate to /quicktype-config.tmp through CooperativeFileWriter.")
number("Verify exact byte count and parse the temporary file back as JSON.")
number("Rename the current configuration to a backup.")
number("Rename the verified temporary file to the active path.")
number("If activation fails, restore the backup; otherwise remove the backup, resume the host stack, re-enable the watchdog, and reset bridge-derived state.")
doc.add_paragraph("At boot, recovery promotes a backup if the active file is absent and deletes stale temporary state. A missing saved configuration deliberately activates the legacy hard-coded keypad mappings.")

page_break()
doc.add_heading("7. Primary firmware execution",1)
doc.add_heading("Laptop-facing composite USB",2)
doc.add_paragraph("Before normal Arduino setup completes, the sketch detaches TinyUSB, adds its final CDC and HID interfaces, and reattaches so the computer does not cache a transient serial-only descriptor. The primary then acts as a keyboard, Consumer Control device, standard mouse, and serial configurator endpoint.")
doc.add_heading("Core division",2)
table(["Execution context","Responsibilities"],[
    ("Core 0: setup()/loop()","Native USB output, status LED, browser serial protocol, telemetry heartbeat, RTC/LittleFS initialization, watchdog servicing, and expansion emission."),
    ("Core 1: setup1()/loop1()","Hardware UART bridge ingestion and selected downstream-host input work; guarded by host-stack suspension during flash writes."),
], [1.7,4.8])
doc.add_paragraph("The design protects the laptop-facing USB connection from downstream stalls. Input processing and output emission use cooperative service points, and the firmware preserves HID FIFO order so a queued key-up cannot overwrite its preceding key-down.")
doc.add_heading("Input arbitration",2)
bullet("Valid framed UART heartbeats mark the bridge active. The UART path carries mount metadata, descriptor chunks, reports, unmounts, logs, and bridge version.")
bullet("If UART activity expires after 1.5 seconds, the primary clears held keys/buttons, discards bridge state, and resumes PIO USB.")
bullet("In PIO mode, boot-capable keyboards switch to boot protocol. Mice remain in report protocol to preserve wheel, pan, and extended buttons.")
bullet("In bridge mode, report descriptors are forwarded so the primary can parse report IDs, keyboard reports, Consumer Control usages, and mouse bit fields.")
doc.add_heading("HID output behavior",2)
doc.add_paragraph("Keyboard reports are either passed through or intercepted by configured rules. Consumer Control reports preserve media keys. Standard mouse reports pass through movement, vertical wheel, horizontal pan, and up to eight buttons; mouse remapping is not implemented. PC keyboard LED state is retained because Windows Alt-code generation may require Num Lock handling.")

doc.add_heading("8. Macro/template engine",1)
doc.add_paragraph("typeExpansionTemplate() walks template text, resolves custom and RTC-backed placeholders, recognizes inline double-brace HID chords, converts selected Unicode bullets to Windows Alt-key sequences, and emits key reports with a rule-specific delay (5 ms by default for new/unspecified rules). Output is cooperative so USB and watchdog work continue during long macros.")
table(["Syntax","Example","Result"],[
    ("Custom placeholder","{initials}","Substitute configured value at execution time."),
    ("Clock field","{date}, {time_24_seconds}","Read DS3231 and format the value."),
    ("Relative/custom date","{date+1}, {date:MM/D/YY}","Offset or format the current RTC date."),
    ("Inline HID key","{{ENTER}}, {{WIN+ALT+K}}","Tap a key or modifier chord within text."),
    ("Structured step","key:CTRL+V","Execute shortcut semantics as one macro step."),
], [1.45,1.85,3.2])
note("Failure behavior","RTC-backed output is skipped when the DS3231 is unavailable or returns an invalid time. Unsupported shortcut tokens fail the action instead of emitting an unintended chord.")

page_break()
doc.add_heading("9. RTC and time metadata",1)
doc.add_paragraph("The DS3231 is connected over I²C on GPIO6 (SDA) and GPIO7 (SCL), address 0x68. The firmware reads/writes its BCD registers directly, checks the oscillator-stop flag, validates calendar values, and stores timezone name/offset separately because the RTC itself contains local calendar time but no timezone database.")
bullet("The website’s Sync Clock command sends the browser’s local date/time, IANA timezone name when available, and numeric UTC offset.")
bullet("Timestamp.txt can initialize the clock at boot; an option can rename it after success.")
bullet("Tokens include date/time variants, weekday/month fields, week/day-of-year/quarter, timezone data, relative dates, and custom format strings.")

doc.add_heading("10. UART bridge protocol",1)
doc.add_paragraph("The bridge protocol is binary and separate from the browser’s JSON protocol. Both firmware directories contain the same QuickTypeUartProtocol.h contract. A packet consists of magic bytes Q,T; version 2; type; rolling sequence byte; payload length; XOR checksum; and up to 72 payload bytes.")
code("'Q' 'T' | version=2 | type | sequence | length | checksum | payload[length]")
table(["Type","Payload purpose"],[
    ("HEARTBEAT","Prove bridge liveness every 250 ms."),
    ("HID_MOUNT","Device address, interface, protocol, total descriptor length."),
    ("HID_DESCRIPTOR","Offset plus a descriptor chunk (60-byte chunk capacity)."),
    ("HID_REPORT","Device/interface prefix followed by raw report bytes."),
    ("HID_UNMOUNT","Release interface state."),
    ("LOG / VERSION","Forward diagnostics and bridge firmware identity."),
], [1.65,4.85])
doc.add_heading("Bridge lifecycle",2)
number("Initialize native TinyUSB host mode and keep all interfaces in report protocol.")
number("On mount, cache up to four interfaces and up to 1,024 descriptor bytes per interface; send mount and descriptor chunks.")
number("Submit interrupt-IN report requests. Each callback forwards the raw report and immediately schedules the next receive.")
number("Every 250 ms send a heartbeat; every 2 seconds re-advertise version, diagnostics, and mounted snapshots so either board can restart independently.")
number("LED state communicates red initialization failure, blue host ready/no device, green mounted, and a cyan flash on reports.")

doc.add_heading("11. Shared expansion sets",1)
doc.add_paragraph("catalog.json names discoverable sets. A quicktype-set/v2 document contains stable set and entry IDs, metadata, typed expansions, keypad expansions, and placeholder definitions. v1 remains importable as typed-only content.")
bullet("Required catalog entries are checked during connection and installed when contents differ.")
bullet("Applying an update replaces only rules managed by that set. Personal rules, unrelated keypad assignments, placeholder values, device name, and color are preserved.")
bullet("Managed rules are read-only during ordinary device editing but can be cloned into personal rules.")
bullet("Content comparison—not revision number alone—detects updates.")
bullet("A device install writes the merged configuration and reads it back before reporting success; a disconnected apply changes browser editor state only.")

page_break()
doc.add_heading("12. Diagnostics and reliability",1)
table(["Tool/sketch","Use"],[
    ("Host_Diagnostic.ino","Minimal reproduction for PIO USB host enumeration and keyboard reports."),
    ("Bridge_Diagnostic.ino","Native USB host diagnostics, descriptor parsing, protocol switching, and report visibility."),
    ("tools/serial_probe.c","Small serial-port connectivity/protocol probe."),
    ("tools/telemetry_monitor.py","POSIX 115200-baud telemetry polling and logging."),
    ("tools/usb_stability_monitor.ps1","Windows long-running serial/USB stability checks with anomaly counters."),
], [2.25,4.25])
doc.add_paragraph("Telemetry exposes protocol activity, configuration writes/failures, serial disconnects, host mount/unmount and report-request failures, loop timing, selected input mode, interface state, firmware/build identity, uptime, and free heap. These signals help distinguish browser transport issues, downstream-host failures, bridge loss, and resource pressure.")
doc.add_heading("Reliability design choices",2)
bullet("The vendored Pico-PIO-USB code is pinned to a documented upstream commit and patched to bound host waits, reset EOP receive state, remove a racy poll, and debounce short SE0 glitches.")
bullet("A pending interrupt-IN transfer while a keyboard is idle is treated as normal and is not aborted on a timer.")
bullet("A downstream host stall must not reboot the RP2040 or deliberately disconnect laptop-facing HID.")
bullet("Transactional configuration files and boot recovery reduce corruption risk across interrupted writes.")
bullet("Bridge snapshots and heartbeats allow independent primary/bridge restarts and deterministic failover.")

doc.add_heading("13. Build and deployment",1)
doc.add_heading("Website",2)
bullet("Serve index.html over HTTPS or localhost in Chrome/Edge; Web Serial is unavailable in many other browser contexts.")
bullet("Publish expansion-sets beside index.html so relative catalog/set URLs resolve.")
bullet("No compilation step is required. Update the quicktype-build timestamp whenever index.html changes.")
doc.add_heading("Primary firmware",2)
bullet("RP2040 board package using the Seeed XIAO RP2040 profile used by this project.")
bullet("USB Stack: Adafruit TinyUSB; Flash: 2 MB with LittleFS (256 KB partition recommended); CPU: 240 MHz.")
bullet("Libraries: Adafruit TinyUSB, Pico-PIO-USB, ArduinoJson 7. A 2 MB/no-FS build cannot persist website configuration.")
doc.add_heading("Bridge firmware",2)
bullet("Build for RP2040-Zero hardware with USB Stack set to Adafruit TinyUSB Host (native).")
bullet("Cross UART TX/RX on GPIO4/GPIO5, connect ground, and provide proper downstream host/OTG VBUS wiring.")
note("Hardware qualification","The 240 MHz primary setting is an RP2040 overclock chosen because PIO USB requires an exact multiple of 120 MHz and the repository notes successful long-host testing at 240 MHz. Device-level qualification remains necessary.")

doc.add_heading("14. End-to-end traces",1)
doc.add_heading("Writing configuration",2)
code("Editor change → browser normalize/validate → compact active rules → size check\n→ Web Serial set-config → firmware parse/compile → temp write/verify\n→ backup swap → activate runtime rules → success response → browser read-back")
doc.add_heading("Executing a typed expansion",2)
code("Downstream HID report → decode key/source → forward printable input + buffer suffix\n→ delimiter arrives → scope/trigger match → backspace trigger text\n→ resolve placeholders/RTC/inline chords → emit HID sequence → optionally restore Space")
doc.add_heading("Bridge failure and failover",2)
code("Bridge heartbeat/reports stop → 1.5 s timeout → release held input\n→ clear bridge interface state → select PIO USB → continue laptop-facing USB service")

page_break()
doc.add_heading("15. Maintenance guide",1)
table(["Change","Likely touch points","Verification focus"],[
    ("New browser command","index.html send/response logic; firmware handleProtocolLine()","Version envelope, timeout/error response, max line size."),
    ("New config field","Browser defaults/normalization/serialization; firmware compile/save/load","Backward compatibility and atomic read-back."),
    ("New trigger behavior","Browser conflict validation; firmware matching/buffering","Delimiter restoration, scope overlap, pass-through."),
    ("New template token","Browser token list/preview; firmware rtcTokenValue()/template engine","Preview parity, invalid RTC behavior, formatting."),
    ("New bridge packet","Both QuickTypeUartProtocol.h copies; sender and parser","Length/checksum, restart snapshots, version compatibility."),
    ("HID support change","Descriptor parser, report routing, laptop HID descriptors","Report IDs, key-up ordering, media/mouse passthrough."),
    ("Shared-set schema change","Catalog/set normalization, merge and export paths","Stable IDs, preservation of personal content, v1 import."),
], [1.35,3.05,2.1])
doc.add_heading("Key invariants",2)
bullet("The two UART protocol headers must remain byte-for-byte compatible.")
bullet("Browser protocol qt and firmware CONFIG_SCHEMA_VERSION must agree.")
bullet("Active rule count must not exceed 128; serialized browser payload must remain below the firmware’s 32 KB ceiling.")
bullet("A configuration must compile successfully before it becomes active or replaces the stored file.")
bullet("A downstream USB problem must not destabilize laptop-facing HID.")
bullet("Managed-set replacement must preserve user-owned rules and device personalization.")

doc.add_heading("16. Source landmarks",1)
table(["Concern","Primary symbol/file"],[
    ("Browser bootstrap","initApp() in index.html"),
    ("Serial request correlation","sendDeviceCommand(), readSerialLoop(), handleSerialLine()"),
    ("Config normalization","normalizeImportedConfiguration(), buildDeviceConfiguration()"),
    ("Shared-set merge","buildConfigurationFromSharedSet(), performSharedSetReplacement()"),
    ("Protocol dispatcher","handleProtocolLine() in RTC_Setup.ino"),
    ("Atomic persistence","saveConfiguration(), recoverConfigurationStorage()"),
    ("Macro execution","executeConfiguredRule(), typeExpansionTemplate(), rtcTokenValue()"),
    ("Main scheduling","setup()/loop() and setup1()/loop1()"),
    ("UART wire contract","QuickTypeUartProtocol.h in both firmware projects"),
    ("Bridge forwarding","writeHidMount(), writeHidReport(), TinyUSB callbacks"),
], [2.1,4.4])
doc.add_paragraph("This guide describes the repository as reviewed. For exact behavior during maintenance, treat the source and its adjacent README files as authoritative, especially where firmware version comments, board-package behavior, or vendored USB-host code evolve.")

doc.save(OUT)
print(OUT.resolve())
